# -*- coding: utf-8 -*-
"""Generate the Microsoft Word document with alt text for the figures and
tables of the main text and the Supplementary Material (AIP Publishing
accessibility requirement).

Run:  py -3.11 results_reviewers/make_alt_text_docx.py
Output: manuscript/Alt_Text_for_Figures_and_Tables.docx
"""
import os

from docx import Document
from docx.shared import Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)

# (kind, label, caption-or-title, alt-text)
MAIN = [
    ("Figure", "Figure 1",
     "Second-mode energy fraction f2 and time-dependent energy fraction ft of the identified control versus Reynolds number for the Fourier and modified-Chebyshev parametrizations.",
     "Two-panel line graph. Left panel: the second-mode energy fraction f2 (in percent) on the vertical axis versus Reynolds number on the horizontal axis for the Fourier basis (solid curve with circles) and the modified-Chebyshev basis (dashed curve with squares). The Fourier curve reaches about 92 percent at Re = 100 and 500 and drops to 64 percent at Re = 1000, while the Chebyshev curve stays near zero. Right panel: the time-dependent energy fraction ft (in percent) versus Reynolds number; the Fourier curve stays near or below 1.5 percent at Re = 100 and 500 and rises to about 32 percent at Re = 1000, while the Chebyshev curve exceeds 93 percent at all three Reynolds numbers."),
    ("Figure", "Figure 2",
     "Reynolds-number dependence of the LBM-realized response to the PINN Fourier control. Left: flow-response amplitude extracted from the mean vertical-centerline velocity. Right: total enstrophy of the tested lid profiles.",
     "Two-panel line graph of the independently realized flow response versus Reynolds number. Left panel: the flow-response amplitude a2 of the mean flow along the vertical centerline (dimensionless) on a logarithmic vertical axis versus Reynolds number (100, 500, 1000); the amplitude increases from about 3e-3 at Re = 100 to 5e-2 at Re = 1000. Right panel: total enstrophy (dimensionless) versus Reynolds number, lower for the controlled profiles than for the uniform lid at every Reynolds number."),
    ("Figure", "Figure 3",
     "Actuation-energy sweep at Re = 500. Left: second-mode energy fraction f2. Right: coefficient A2 of the stationary second spatial harmonic.",
     "Two-panel graph, both axes with log-scaled target energy E* from 0.01 to 2 on the horizontal axis. Left panel: second-mode energy fraction f2 (in percent) versus target energy; the fraction stays near 92 percent for targets up to 0.25, then decreases to about 74 percent at E* = 2. Right panel: coefficient A2 (dimensionless) versus target energy; A2 rises monotonically from about 0.32 at the smallest target to about 1.69 at E* = 2."),
    ("Figure", "Figure 4",
     "Aspect-ratio dependence at Re = 500, E* = 0.25. Left: identified Fourier coefficients A1, A2, and A3. Right: second-mode fraction f2 and time-dependent fraction ft.",
     "Two-panel graph comparing the square (aspect ratio 1, label '1') and the 2:1 rectangular (aspect ratio 2) cavities. Left panel: grouped-bar chart of the Fourier coefficients A1, A2, A3 (dimensionless); A2 is close to 0.68 in both geometries while A1 and A3 are near zero. Right panel: grouped-bar chart of the second-mode fraction f2 (about 92 percent for both geometries) and the time-dependent fraction ft (1.3-1.6 percent)."),
    ("Figure", "Figure 5",
     "Identified lid controls U_lid(s), with s = x/Lx, at Re = 500 for the square and 2:1 rectangular cavities.",
     "Two-panel line graph of the identified lid velocity profile U_lid as a function of the normalized horizontal coordinate s = x/Lx, for the square cavity (left) and the 2:1 rectangular cavity (right). Each panel shows superposed curves: the reference A2 sin(2 pi s), the complete PINN control, and its reduced-mode reconstruction; the three curves nearly coincide and follow a positive lobe near the left wall and a negative lobe near the right wall."),
    ("Figure", "Figure 6",
     "Verification of the independent LBM-MRT solver against Ghia et al. for the uniform-lid cavity at Re = 100 and Re = 1000.",
     "Two-panel line graph of the u-velocity profile along the vertical centerline (x = 0.5) as a function of the vertical coordinate y, for Re = 100 (left) and Re = 1000 (right). The LBM solution and the Ghia reference data are plotted on the same axes and agree to within a few percent, validating the uncontrolled solver."),
    ("Figure", "Figure 7",
     "Three-grid convergence study at Re = 500 for the stationary second-harmonic control.",
     "Two-panel graph of the grid-convergence study at grids of 128, 256, and 512 points. Left panel: integrated viscous dissipation against grid size, increasing slightly from 0.0094 to 0.0109 as the grid is refined toward the extrapolated value near 0.0114. Right panel: deviation of the dissipation from the extrapolated value versus grid spacing on log axes, with an observed convergence order p = 1.08."),
    ("Figure", "Figure 8",
     "Integrated dissipation of the LBM-realized flows under the uniform lid, first spatial harmonic, stationary second harmonic, and complete PINN mean control.",
     "Grouped-bar chart with one bar group per Reynolds number (100, 500, 1000). Each group shows the integrated dissipation of the uniform-lid, sin(pi x), A2 sin(2 pi x), and PINN mean controls; the controlled cases are substantially lower than the uniform reference at every Reynolds number, with relative reductions between about 73 and 80 percent."),
    ("Figure", "Figure 9",
     "Consolidated identification and independent-verification results.",
     "Four-panel summary graph. Panel (a): integrated dissipation versus Reynolds number for the stationary second-harmonic control, the complete PINN mean control, and the uniform reference, showing large reductions for the controlled cases. Panel (b): boundary-control coefficient a2 (control) and independently extracted flow-response amplitude a2 (flow) versus Reynolds number, showing a compatible trend with very different magnitudes. Panel (c): fluctuating-to-mean-flow kinetic-energy ratio RK (in percent) versus Reynolds number for the time-dependent Fourier and modified-Chebyshev branches, together with the adopted 1 percent threshold; the Fourier branch is below the threshold, the Chebyshev branch far above it. Panel (d): dominant modal indices of the mean and fluctuation fields for the two branches."),
    ("Table", "Table I",
     "Baseline Fourier branch. Er is the realized control energy, A2 the coefficient of the stationary second spatial harmonic, f2 its energy fraction, and ft the time-dependent energy fraction.",
     "Table with one row per Reynolds number (100, 500, 1000) and four columns: realized control energy Er, coefficient A2, second-mode fraction f2 (in percent), and time-dependent fraction ft (in percent). At Re = 100 and 500 the second-mode fraction is 91.6 and 92.2 percent with ft below 1.3 percent, while at Re = 1000 it drops to 64.0 percent with ft = 32.2 percent."),
    ("Table", "Table II",
     "Actuation-energy sweep at Re = 500. E* is the prescribed dimensionless target, Er the realized dimensionless energy, f2 the second-mode fraction, and ft the time-dependent fraction.",
     "Table of seven rows, one per target energy E* from 0.01 to 2.00, with columns E*, realized energy Er, second-mode fraction f2 (in percent), and time-dependent fraction ft (in percent). The second-mode fraction stays between 92.1 and 92.8 percent for targets up to 0.25 and decreases to 74.0 percent at E* = 2.00, while ft rises from 0.4 to 21.8 percent."),
    ("Table", "Table III",
     "Ten-seed training study at Re = 500. Er is realized control energy, f2 the second-mode fraction, ft the time-dependent fraction, and L_total the final total loss.",
     "Table of ten rows, one per independent training seed (0 to 9), with columns Er, f2 (in percent), ft (in percent), and final total loss. The realized energies lie in 0.240-0.268 and the final losses in 1.371-1.394, while the modal composition splits into second-mode-dominated runs (seeds 1, 2, 6, 7), temporally dominated runs (seeds 3, 5, 8, 9), and two intermediate runs (seeds 0, 4)."),
    ("Table", "Table IV",
     "Loss-ablation study at Re = 500. Er is realized control energy, f2 the second-mode fraction, and ft the time-dependent fraction.",
     "Table with eight configuration rows and three columns (Er, f2 in percent, ft in percent). The baseline gives f2 = 92.2 percent; removing the variance regularization L_var redirects the solution to a first-harmonic branch with f2 near 0.2 percent (whether or not L_Re or L_diss is also removed), while halving or doubling the variance weight preserves the branch with f2 between 89.3 and 92.8 percent."),
    ("Table", "Table V",
     "Network-capacity and control-space truncation study at Re = 500.",
     "Two-part table. Left: the three network architectures (small, baseline, large) with the range of f2 (in percent) obtained over three seeds (90.3-93.6, 89.2-91.1 excluding the intermediate branch, 86.7-92.2). Right: control-space truncations (Nx,Nt) of (4,3), (6,1), (6,3), (6,5), (8,5), (8,9), corresponding to 12 to 72 coefficients, with Er and f2 in percent; f2 ranges from 87.3 to 92.3 percent."),
    ("Table", "Table VI",
     "Aspect-ratio dependence at Re = 500. A2 is the stationary second-harmonic coefficient, Er the realized control energy, f2 the second-mode fraction, and ft the time-dependent fraction.",
     "Table of two rows (aspect ratios Lx/Ly = 1 and 2) with columns A2, Er, f2 (in percent), and ft (in percent). A2 is 0.6835 (square) and 0.6770 (2:1), Er is 0.2534 and 0.2514, f2 is 92.2 and 91.2 percent, and ft is 1.3 and 1.6 percent."),
    ("Table", "Table VII",
     "Comparison of the Fourier and modified-Chebyshev control parametrizations.",
     "Table of two block rows (Fourier and modified-Chebyshev bases), each with three Reynolds-number rows (100, 500, 1000), and columns Er, f2 (in percent), ft (in percent), and reconstruction RMSE. The Fourier basis gives f2 near 92 percent at Re = 100-500 and 64.0 percent at Re = 1000 with machine-precision reconstruction, whereas the modified-Chebyshev basis gives ft between 93 and 99 percent with reconstruction RMSE of order 1e-3."),
    ("Table", "Table VIII",
     "Centerline errors of the independent LBM-MRT solver relative to the Ghia benchmark.",
     "Table of two rows (Re = 100 and 1000) with the L2 and L-infinity errors of the u- and v-velocity components along the centerline. At Re = 100 the L2(u) error is 0.0051 and at Re = 1000 it is 0.0141, with the velocity-component errors all at the few-percent level."),
    ("Table", "Table IX",
     "Integrated dissipation of the steady LBM realizations for the uniform lid, first harmonic, stationary second harmonic, and time-mean PINN control at three Reynolds numbers.",
     "Table of three rows (Re = 100, 500, 1000) and four columns of integrated dissipation: uniform lid, sin(pi x), A2 sin(2 pi x), and PINN (mean). The stationary second-harmonic control gives the lowest dissipation at every Reynolds number (0.0345, 0.0104, 0.0047), and the PINN mean control is close to it (0.0378, 0.0119, 0.0052)."),
    ("Table", "Table X",
     "Time-dependent LBM assessment of the Fourier and modified-Chebyshev branches. RK = K_fluct/K_mean is reported as a percentage.",
     "Table of six rows (Fourier and Chebyshev branches at Re = 100, 500, 1000) with columns RK (in percent) and pass/fail criterion. The Fourier branch passes the adopted 1 percent threshold at every Reynolds number (0.15, 0.04, 0.65 percent); the modified-Chebyshev branch fails by large margins (847, 44.0, 125 percent)."),
    ("Table", "Table XI",
     "Leave-one-condition-out evaluation of the symbolic representation on the auxiliary nine-condition dataset.",
     "Table of one row of R2 values (coefficient of determination, dimensionless) over the nine held-out Reynolds conditions 100 to 1000; R2 ranges from 0.9326 to 0.9908, showing that a compact symbolic expression reproduces the identified control across the sampled conditions."),
]

SI = [
    ("Table", "Table S1",
     "Network architecture sweep at Re = 500, Fourier basis, three seeds per architecture.",
     "Table of three rows (small, baseline, large architectures) with columns: number of trainable parameters (8774, 44824, 139624), realized-energy range Er, second-mode fraction range f2 (in percent), and temporal fraction range ft (in percent). All three capacities recover the second-mode-dominated branch with f2 in the high-80s to low-90s percent range."),
    ("Table", "Table S2",
     "Kinetic-energy diagnostics of the LBM realizations under the time-dependent Fourier (PINN) and modified-Chebyshev controls.",
     "Table of six rows (PINN Fourier and Chebyshev time-dependent controls at Re = 100, 500, 1000) with columns K_mean, K_fluct, and RK (in percent). The Fourier branch has RK = 0.15, 0.04, 0.65 percent; the Chebyshev branch has RK = 847, 44.0, and 125 percent, dominated by its small mean-flow kinetic energy at Re = 100."),
    ("Figure", "Figure S1",
     "Velocity fields of the Fourier branch at Re = 100, 500, and 1000 realized in the LBM solver.",
     "Composite of three velocity-magnitude fields of the LBM-realized controlled flow, one per Reynolds number (100, 500, 1000). Each field shows a central circulation set up by the dominant second spatial mode, with low velocities near the walls."),
    ("Figure", "Figure S2",
     "Viscous dissipation fields of the Fourier branch at Re = 100, 500, and 1000 realized in the LBM solver.",
     "Composite of three viscous-dissipation-density fields (dimensionless), one per Reynolds number. The dissipation is concentrated in the shear layers set up by the forced lid motion, thinning as the Reynolds number increases."),
    ("Figure", "Figure S3",
     "Modal analysis of the vertical-centerline velocity u(0.5,y) of the LBM-realized flows under the PINN (mean) and A2 sin(2 pi x) lid profiles.",
     "Composite graph of the Fourier modal content of the vertical-centerline velocity of the LBM-realized mean flows under the complete PINN control and its dominant-mode truncation, at Re = 100, 500, and 1000. The mean flow of the Fourier branch is dominated by the second spatial contribution at every Reynolds number."),
    ("Figure", "Figure S4",
     "Centerline velocity profiles u(0.5,y) and v(x,0.5) of the LBM-realized flows at Re = 100, 500, and 1000.",
     "Six-panel composite: for each Reynolds number, the u-velocity profile along the vertical centerline and the v-velocity profile along the horizontal centerline, under the uniform, sin(pi x), A2 sin(2 pi x), and PINN mean lid profiles. The controlled profiles depart strongly from the uniform-lid reference."),
    ("Figure", "Figure S5",
     "Modal spectra of the recovered control at Re = 100.",
     "Two-panel bar/line graph of the spatial and temporal modal spectra of the recovered control at Re = 100. The spatial spectrum is dominated by the second spatial mode; the temporal spectrum is essentially flat and small."),
    ("Figure", "Figure S6",
     "Modal spectra of the recovered control at Re = 500.",
     "Two-panel bar/line graph of the spatial and temporal modal spectra of the recovered control at Re = 500. The spatial spectrum peaks at the second spatial mode with a small temporal content."),
    ("Figure", "Figure S7",
     "Modal spectra of the recovered control at Re = 1000.",
     "Two-panel bar/line graph of the spatial and temporal modal spectra of the recovered control at Re = 1000. The second spatial mode remains the largest individual contribution, while the temporal content is more pronounced than at lower Reynolds numbers."),
    ("Figure", "Figure S8",
     "Temporal response of the LBM realization to the time-dependent Fourier (PINN) control at Re = 100.",
     "Temporal-response panel set at Re = 100: mean and fluctuating fields, centerline profiles, a probe time series over eight forcing periods, its frequency spectrum, and eight-phase snapshots. The fluctuation amplitude remains below 1 percent of the mean-flow kinetic energy (quasi-stationary)."),
    ("Figure", "Figure S9",
     "Temporal response of the LBM realization to the time-dependent Fourier (PINN) control at Re = 500.",
     "Temporal-response panel set at Re = 500 analogous to Figure S8; the realization is quasi-stationary, with the fluctuation-to-mean kinetic ratio RK = 0.04 percent."),
    ("Figure", "Figure S10",
     "Temporal response of the LBM realization to the time-dependent Fourier (PINN) control at Re = 1000.",
     "Temporal-response panel set at Re = 1000 analogous to Figure S8; the realization remains quasi-stationary with RK = 0.65 percent."),
    ("Figure", "Figure S11",
     "Temporal response of the LBM realization to the time-dependent modified-Chebyshev control at Re = 100.",
     "Temporal-response panel set at Re = 100 analogous to Figure S8 but for the modified-Chebyshev control: the fluctuations dominate the kinetic energy (RK = 847 percent), driven by the very small mean-flow kinetic energy."),
    ("Figure", "Figure S12",
     "Temporal response of the LBM realization to the time-dependent modified-Chebyshev control at Re = 500.",
     "Temporal-response panel set at Re = 500 for the modified-Chebyshev control; fluctuations dominate the kinetic energy with RK = 44 percent."),
    ("Figure", "Figure S13",
     "Temporal response of the LBM realization to the time-dependent modified-Chebyshev control at Re = 1000.",
     "Temporal-response panel set at Re = 1000 for the modified-Chebyshev control; fluctuations dominate the kinetic energy with RK = 125 percent."),
    ("Figure", "Figure S14",
     "Global temporal response of the LBM realizations under all tested lid controls at Re = 100, 500, and 1000.",
     "Composite figures collecting the time series and fluctuation content of the LBM realizations under the uniform, sin(pi x), A2 sin(2 pi x), PINN mean, and time-dependent Fourier and modified-Chebyshev profiles at the three Reynolds numbers. The Fourier-related profiles settle into quasi-stationary states while the modified-Chebyshev profile remains strongly time-dependent."),
]


def add_block(doc, kind, label, title, alt):
    h = doc.add_heading(f"{label} ({kind.lower()})", level=2)
    doc.add_paragraph().add_run("Caption/Title: ").bold = True
    p = doc.add_paragraph(title)
    doc.add_paragraph().add_run("Alt text: ").bold = True
    doc.add_paragraph(alt)


def main():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Calibri"
    st.font.size = Pt(11)

    doc.add_heading("Alt text for figures and tables", level=0)
    doc.add_paragraph(
        "Physics of Fluids manuscript PF# POF26-AR-10710: "
        "\u201cPhysics-informed identification and independent numerical "
        "verification of conditional modal collapse in active control of "
        "lid-driven cavity flow.\u201d "
        "This document provides the alt text for every figure and table of "
        "the main text and of the Supplementary Material, as requested by "
        "AIP Publishing for accessibility."
    )

    doc.add_heading("Main text", level=1)
    for kind, label, title, alt in MAIN:
        add_block(doc, kind, label, title, alt)

    doc.add_heading("Supplementary Material", level=1)
    for kind, label, title, alt in SI:
        add_block(doc, kind, label, title, alt)

    out = os.path.join(ROOT, "manuscript", "Alt_Text_for_Figures_and_Tables.docx")
    doc.save(out)
    print(f"-> {out}")


if __name__ == "__main__":
    main()